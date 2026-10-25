"""Lista leków refundowanych – obwieszczenie Ministra Zdrowia (gov.pl, XLSX, co 2–3 miesiące).

    python -m app.refund            # importuje data/raw/refundacja.xlsx (pobiera, jeśli brak)
    python -m app.refund --pobierz  # zawsze pobiera najnowsze obwieszczenie

Z załącznika bierzemy arkusz A1 (leki w aptece na receptę: cena, limit, poziom odpłatności, dopłata
pacjenta – osobno dla każdego zakresu wskazań) oraz listy bezpłatnych leków: D1 (seniorzy 65+, „S”),
D2 (dzieci i młodzież do 18 lat, „DZ”) i E (kobiety w ciąży, „C”). Łączymy po kodzie EAN (GTIN).
"""
import argparse
import html
import json
import os
import re
import urllib.request
from datetime import datetime
from pathlib import Path

from .importer import DATA_DIR

LIST_URL = "https://www.gov.pl/web/zdrowie/obwieszczenia-ministra-zdrowia-lista-lekow-refundowanych"
BASE_URL = "https://www.gov.pl"
RAW_XLSX = DATA_DIR / "raw" / "refundacja.xlsx"
RAW_META = DATA_DIR / "raw" / "refundacja.meta.json"
JSON_PATH = DATA_DIR / "refundacja.json"

FREE_SHEETS = {"D1": "S", "D2": "DZ", "E": "C"}
FREE_LABELS = {"S": "bezpłatny dla seniorów 65+", "DZ": "bezpłatny dla dzieci i młodzieży do 18 lat",
               "C": "bezpłatny dla kobiet w ciąży"}
_ANNOUNCEMENT = re.compile(r'href="(/web/zdrowie/obwieszczenie-ministra-zdrowia-z-dnia-[^"]+)"[^>]*>\s*([^<]+)')
_ATTACHMENT = re.compile(r'href="(/attachment/[0-9a-f-]+)"[^>]*aria-label="([^"]*)"', re.S)


def _get(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (farmacja-pomocnik/1.0)"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return resp.read()


def find_latest() -> tuple[str, str, str]:
    """(tytuł obwieszczenia, adres strony, adres załącznika XLSX) – najnowsze obwieszczenie główne."""
    page = _get(LIST_URL).decode("utf-8", "replace")
    for path, title in _ANNOUNCEMENT.findall(page):
        title = " ".join(html.unescape(title).split())
        detail = _get(BASE_URL + path).decode("utf-8", "replace")
        for att, label in _ATTACHMENT.findall(detail):
            label = " ".join(html.unescape(label).split()).lower()
            if "zalacznik do obwieszczenia" in label and "xlsx" in label:
                return title, BASE_URL + path, BASE_URL + att
    raise RuntimeError("Nie znaleziono załącznika XLSX z listą refundacyjną na stronie gov.pl")


def download() -> dict:
    title, page_url, xlsx_url = find_latest()
    RAW_XLSX.parent.mkdir(parents=True, exist_ok=True)
    tmp = RAW_XLSX.with_suffix(".xlsx.part")
    tmp.write_bytes(_get(xlsx_url))
    os.replace(tmp, RAW_XLSX)
    meta = {"tytul": title, "strona": page_url, "plik": xlsx_url}
    RAW_META.write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
    return meta


def _gtin(value) -> str:
    return re.sub(r"\D", "", str(value or "")).lstrip("0")


def _text(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def import_xlsx(xlsx_path: Path = RAW_XLSX, json_path: Path = JSON_PATH) -> int:
    import openpyxl  # tylko do importu; aplikacja czyta gotowy JSON

    meta = json.loads(RAW_META.read_text(encoding="utf-8")) if RAW_META.exists() else {}
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    packages: dict[str, dict] = {}
    for row in list(wb["A1"].iter_rows(values_only=True))[2:]:
        gtin = _gtin(row[4])
        if not gtin:
            continue
        entry = packages.setdefault(gtin, {
            "nazwa": _text(row[2]), "opakowanie": _text(row[3]), "grupa_limitowa": _text(row[7]),
            "cena": _text(row[11]), "limit": _text(row[12]), "wskazania": [], "bezplatny": {},
        })
        indication = _text(row[13])
        if row[14]:
            indication += f" (pozarejestracyjnie: {_text(row[14])})"
        entry["wskazania"].append({"zakres": indication, "odplatnosc": _text(row[15]), "doplata": _text(row[16])})
    for sheet, code in FREE_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        for row in list(wb[sheet].iter_rows(values_only=True))[2:]:
            gtin = _gtin(row[4] if len(row) > 4 else "")
            if gtin in packages:
                packages[gtin]["bezplatny"][code] = _text(row[5]) if len(row) > 5 else ""
    wb.close()
    data = {"obwieszczenie": meta.get("tytul", ""), "strona": meta.get("strona", ""),
            "zaimportowano": datetime.now().isoformat(timespec="seconds"), "opakowania": packages}
    tmp = json_path.with_suffix(".json.part")
    tmp.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    os.replace(tmp, json_path)
    return len(packages)


def _amount(text: str) -> float | None:
    try:
        return float(text.replace(" ", "").replace(",", "."))
    except (AttributeError, ValueError):
        return None


class RefundDB:
    """Refundacja wg kodu EAN opakowania. Brak pliku = brak danych (aplikacja działa dalej)."""

    def __init__(self, path: Path = JSON_PATH):
        data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
        self.title = data.get("obwieszczenie", "")
        self.page = data.get("strona", "")
        self.packages: dict[str, dict] = data.get("opakowania", {})

    @property
    def loaded(self) -> bool:
        return bool(self.packages)

    @property
    def summary(self) -> dict:
        return {"obwieszczenie": self.title, "strona": self.page, "liczba_opakowan": len(self.packages)}

    def for_gtin(self, gtin: str) -> dict | None:
        return self.packages.get(_gtin(gtin))

    def product_badge(self, product) -> dict | None:
        """Skrót dla listy wyników: najniższa dopłata pacjenta i listy bezpłatne (S / DZ / C)."""
        found = [r for pk in product.packages if not pk["deleted"] and (r := self.for_gtin(pk["gtin"]))]
        if not found:
            return None
        payments = [a for r in found for w in r["wskazania"] if (a := _amount(w["doplata"])) is not None]
        free = sorted({code for r in found for code in r["bezplatny"]}, key=list(FREE_LABELS).index)
        return {"min_doplata": min(payments) if payments else None, "bezplatny": free}


def main() -> None:
    parser = argparse.ArgumentParser(description="Import listy leków refundowanych (obwieszczenie MZ)")
    parser.add_argument("--pobierz", action="store_true", help="pobierz najnowsze obwieszczenie z gov.pl")
    args = parser.parse_args()
    if args.pobierz or not RAW_XLSX.exists():
        print("Pobieranie obwieszczenia…")
        print(download()["tytul"])
    print(f"Zaimportowano {import_xlsx()} opakowań refundowanych do {JSON_PATH}")


if __name__ == "__main__":
    main()
